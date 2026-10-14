import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scrape_chrysler2 import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    is_product_detail_page,
    load_results,
    normalize_part_number,
    parse_msrp,
    prepare_crawl_rows,
    read_oem_values,
    save_results,
)


class ChryslerPartsGiantTests(unittest.TestCase):
    def test_normalize_and_parse_msrp(self):
        self.assertEqual(normalize_part_number("0508-3285-AA"), "05083285AA")
        self.assertEqual(parse_msrp("Price $16.31\nMSRP: $23.35"), 23.35)
        self.assertIsNone(parse_msrp("Price: $16.31"))

    def test_moparpartsgiant_product_url_matches_oem_without_leading_zero(self):
        url = "https://www.moparpartsgiant.com/parts/mopar-filter-engine-oil~5083285aa.html"
        self.assertTrue(detail_url_matches_oem(url, "05083285AA"))
        self.assertFalse(detail_url_matches_oem(url, "05083285AB"))

    def test_product_detail_route_is_detected(self):
        page = Mock()
        page.url = "https://www.moparpartsgiant.com/parts/mopar-filter-engine-oil~5083285aa.html"
        self.assertTrue(is_product_detail_page(page))
        page.url = "https://www.moparpartsgiant.com/category/oil-filters.html"
        self.assertFalse(is_product_detail_page(page))

    def test_status_reports_replacement_part(self):
        self.assertEqual(
            detail_status(
                "Mopar 5083285AA Engine Oil Filter",
                "68086920AA",
                23.35,
                "https://www.moparpartsgiant.com/parts/mopar-filter-engine-oil~5083285aa.html",
            ),
            "success_superseded: 5083285aa",
        )

    def test_direct_product_navigation_extracts_msrp(self):
        page = Mock()
        page.url = "https://www.moparpartsgiant.com/"
        search_box = Mock()
        search_box.press.side_effect = lambda *_: setattr(
            page,
            "url",
            "https://www.moparpartsgiant.com/parts/mopar-filter-engine-oil~5083285aa.html",
        )
        page.locator.return_value.inner_text.return_value = "Mopar 5083285AA Engine Oil Filter\nMSRP: $23.35"
        with (
            patch("scrape_chrysler2.dismiss_blocking_popups"),
            patch("scrape_chrysler2.find_search_box", return_value=search_box),
            patch("scrape_chrysler2.wait_for_human_verification"),
            patch("scrape_chrysler2.wait_for_product_msrp", return_value="MSRP: $23.35"),
        ):
            result = crawl_msrp(page, "05083285AA", [])
        self.assertEqual(result, (23.35, "success"))
        search_box.fill.assert_called_once_with("05083285AA")
        search_box.press.assert_called_once_with("Enter")

    def test_excel_read_and_output_round_trip(self):
        with TemporaryDirectory() as directory:
            input_path = Path(directory) / "parts.xlsx"
            workbook = Workbook()
            workbook.active.append(("OEM Number",))
            workbook.active.append(("05083285AA",))
            workbook.save(input_path)
            workbook.close()
            self.assertEqual(read_oem_values(input_path, None, "OEM Number", 1), ["05083285AA"])

            output = Path(directory) / "results.xlsx"
            rows = [("05083285AA", 23.35, date(2026, 10, 12), "success")]
            save_results(output, rows)
            self.assertEqual(load_results(output)[0][:2], rows[0][:2])
            saved = load_workbook(output, data_only=True)
            self.assertEqual(tuple(cell.value for cell in saved.active[1]), ("OEM Number", "MSRP", "Date", "Status"))
            saved.close()

    def test_resume_skips_success_and_retries_failure(self):
        saved = [
            ("05083285AA", 23.35, date(2026, 10, 12), "success"),
            ("05083285AB", None, date(2026, 10, 12), "not_found"),
        ]
        _, pending = prepare_crawl_rows(["05083285AA", "05083285AB"], saved)
        self.assertEqual(pending, [(1, "05083285AB")])


if __name__ == "__main__":
    unittest.main()
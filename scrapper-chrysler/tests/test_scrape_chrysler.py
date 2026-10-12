import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scrape_chrysler import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    find_matching_result,
    is_matching_detail_page,
    load_results,
    normalize_part_number,
    parse_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
)


class ChryslerCrawlerTests(unittest.TestCase):
    def test_normalize_and_parse_msrp(self):
        self.assertEqual(normalize_part_number("0508-3285-AA"), "05083285AA")
        self.assertEqual(parse_msrp("Sale Price: $17.30\nMSRP: $23.35"), 23.35)
        self.assertIsNone(parse_msrp("Price: $17.30"))

    def test_mopar_oem_route_matches_part_number(self):
        self.assertTrue(detail_url_matches_oem("https://www.mymoparparts.com/oem-parts/mopar-engine-oil-filter-5083285aa", "05083285AA"))
        self.assertFalse(detail_url_matches_oem("https://example.com/oem-parts/mopar-engine-oil-filter-5083285aa", "05083285AA"))

    def test_exact_result_is_preferred(self):
        replacement = Mock()
        replacement.is_visible.return_value = True
        replacement.inner_text.return_value = "05083285AA"
        replacement.get_attribute.return_value = "/oem-parts/mopar-engine-oil-filter-5083285aa"
        near_match = Mock()
        near_match.is_visible.return_value = True
        near_match.inner_text.return_value = "05083285AB"
        near_match.get_attribute.return_value = "/oem-parts/mopar-filter-5083285ab"
        links = Mock()
        links.count.return_value = 2
        links.nth.side_effect = lambda index: (near_match, replacement)[index]
        page = Mock()
        page.url = "https://www.mymoparparts.com/search?search_str=05083285AA"
        page.locator.return_value = links
        self.assertIs(find_matching_result(page, "05083285AA"), replacement)

    def test_closest_result_is_selected_when_multiple_results(self):
        closest = Mock()
        closest.is_visible.return_value = True
        closest.inner_text.return_value = "63-21-7-888-850"
        closest.get_attribute.return_value = "/oem-parts/mopar-back-up-light-63217888850"
        farther = Mock()
        farther.is_visible.return_value = True
        farther.inner_text.return_value = "63-21-7-888-852"
        farther.get_attribute.return_value = "/oem-parts/mopar-rear-light-63217888852"
        links = Mock()
        links.count.return_value = 2
        links.nth.side_effect = lambda index: (closest, farther)[index]
        page = Mock()
        page.url = "https://www.mymoparparts.com/search?search_str=6321788885"
        page.locator.return_value = links
        self.assertIs(find_matching_result(page, "6321788885"), closest)

    def test_detail_page_confirms_replaced_oem(self):
        page = Mock()
        page.url = "https://www.mymoparparts.com/oem-parts/mopar-engine-oil-filter-5083285aa"
        body = Mock()
        body.inner_text.return_value = (
            "Engine Oil Filter - MOPAR (05083285AA)\nMSRP: $23.35\nReplaces: 68086920aa"
        )
        page.locator.return_value = body
        self.assertTrue(is_matching_detail_page(page, "05083285AA"))
        self.assertEqual(detail_status("Part Number: 05083285AA", "05083285AA", 23.35), "success")
        self.assertEqual(
            detail_status("Part Number: 05083285AA", "68086920AA", 23.35),
            "success_superseded: 05083285AA",
        )

    def test_crawl_clicks_selected_result_and_reads_msrp(self):
        page = Mock()
        page.url = "https://www.mymoparparts.com/search?search_str=05083285AA"
        search_box = Mock()
        result_link = Mock()
        result_link.get_attribute.return_value = "/oem-parts/mopar-engine-oil-filter-5083285aa"
        result_link.click.side_effect = lambda **_: setattr(
            page, "url", "https://www.mymoparparts.com/oem-parts/mopar-engine-oil-filter-5083285aa"
        )
        page.locator.return_value.inner_text.return_value = "Part Number: 05083285AA\nMSRP: $23.35"
        with (
            patch("scrape_chrysler.find_search_box", return_value=search_box),
            patch("scrape_chrysler.wait_for_human_verification"),
            patch("scrape_chrysler.is_matching_detail_page", return_value=False),
            patch("scrape_chrysler.wait_for_matching_result", return_value=result_link),
        ):
            self.assertEqual(crawl_msrp(page, "05083285AA", []), (23.35, "success"))
        result_link.click.assert_called_once_with(timeout=10_000)

    def test_excel_helpers_and_results_round_trip(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            workbook = Workbook()
            workbook.active.append(("OEM Number",))
            workbook.active.append(("05083285AA",))
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_oem_values(path, None, "OEM Number", 1), ["05083285AA"])
            source = load_workbook(path)
            self.assertEqual(resolve_column(source.active, "A", 1), (1, 2))
            source.close()

            output = Path(directory) / "results.xlsx"
            rows = [("05083285AA", 23.35, date(2026, 10, 12), "success")]
            save_results(output, rows)
            self.assertEqual(load_results(output)[0][:2], rows[0][:2])
            saved = load_workbook(output, data_only=True)
            self.assertEqual(tuple(cell.value for cell in saved.active[1]), ("OEM Number", "MSRP", "Date", "Status"))
            saved.close()

    def test_resume_retries_failed_rows(self):
        saved = [("05083285AA", None, date(2026, 10, 12), "not_found")]
        rows, pending = prepare_crawl_rows(["05083285AA"], saved)
        self.assertEqual(pending, [(0, "05083285AA")])
        self.assertEqual(rows[0][3], "not_found")


if __name__ == "__main__":
    unittest.main()
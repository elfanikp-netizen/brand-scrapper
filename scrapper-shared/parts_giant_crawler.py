from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import time
import urllib.request
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string
from playwright.sync_api import Error as PlaywrightError, Locator, Page, Request, TimeoutError as PlaywrightTimeoutError, sync_playwright


OUTPUT_HEADERS = ("OEM Number", "MSRP", "Date", "Status")
SEARCH_INPUT_SELECTOR = 'input[placeholder="Search by Part Number, Part Name, or Vehicle"]'
MSRP_PATTERN = re.compile(r"\bMSRP\b\s*(?:\([^)]*\))?\s*:?\s*\$?\s*([0-9][0-9,]*(?:\.\d{1,2})?)", re.IGNORECASE)
PART_NUMBER_PATTERN = re.compile(r"(?i)\b(?:Manufacturer\s+)?Part\s*(?:Number|#)\s*:?\s*([A-Z0-9-]*\d[A-Z0-9-]*)")
REPLACES_PATTERN = re.compile(r"(?im)^[ \t]*Replaces[ \t]*:?[ \t]*(?:\r?\n[ \t]*)?([A-Z0-9][A-Z0-9,; -]*)")


@dataclass(frozen=True)
class CrawlerConfig:
    brand: str
    site_url: str
    cdp_port: int
    profile_name: str
    part_label_pattern: str


class HumanVerificationError(RuntimeError):
    pass


def normalize_part_number(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def part_numbers_match(first: object, second: object) -> bool:
    first_normalized = normalize_part_number(first)
    second_normalized = normalize_part_number(second)
    return first_normalized == second_normalized or first_normalized.lstrip("0") == second_normalized.lstrip("0")


def parse_msrp(text: str) -> float | None:
    match = MSRP_PATTERN.search(text)
    return float(match.group(1).replace(",", "")) if match else None


def resolve_column(sheet, column: str, header_row: int) -> tuple[int, int]:
    if header_row > 0:
        for index, cell in enumerate(sheet[header_row], start=1):
            if str(cell.value or "").strip().casefold() == column.strip().casefold():
                return index, header_row + 1
    if column.isdecimal():
        column_index = int(column)
        if column_index < 1:
            raise ValueError("A numeric --column value must be 1 or greater.")
    else:
        try:
            column_index = column_index_from_string(column.strip().upper())
        except ValueError as error:
            raise ValueError(f"Column {column!r} was not found. Use a header, Excel letters, or a 1-based number.") from error
    return column_index, header_row + 1 if header_row > 0 else 1


def read_oem_values(input_path: Path, sheet_name: str | None, column: str, header_row: int) -> list[object]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} was not found in {input_path.name}.")
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        column_index, first_data_row = resolve_column(sheet, column, header_row)
        return [
            row[0]
            for row in sheet.iter_rows(min_row=first_data_row, min_col=column_index, max_col=column_index, values_only=True)
            if row[0] is not None and str(row[0]).strip()
        ]
    finally:
        workbook.close()


def find_search_box(page: Page) -> Locator | None:
    boxes = page.locator(SEARCH_INPUT_SELECTOR)
    for index in range(boxes.count()):
        box = boxes.nth(index)
        if box.is_visible() and box.is_enabled():
            return box
    return None


def dismiss_blocking_popups(page: Page) -> None:
    close_buttons = page.locator(
        'button[aria-label*="close" i], button[title*="close" i], '
        '[role="dialog"] button[class*="close" i], button:has-text("Close"), '
        'button:has-text("Not now"), button:has-text("No thanks")'
    )
    for _ in range(5):
        dismissed = False
        for index in range(close_buttons.count()):
            button = close_buttons.nth(index)
            try:
                if button.is_visible() and button.is_enabled():
                    button.click(timeout=1_000)
                    dismissed = True
                    break
            except PlaywrightError:
                continue
        if not dismissed:
            break
    try:
        page.keyboard.press("Escape")
    except PlaywrightError:
        pass


def page_needs_human_verification(page: Page) -> bool:
    try:
        title = page.title().casefold()
        body = page.locator("body").inner_text(timeout=5_000).casefold()
    except PlaywrightError:
        return False
    indicators = ("performing security verification", "security verification", "verify you are human", "checking your browser")
    return "just a moment" in title or any(indicator in body for indicator in indicators)


def is_site_url(url: str, config: CrawlerConfig) -> bool:
    return urlparse(url).hostname == urlparse(config.site_url).hostname


def wait_for_human_verification(page: Page, challenge_failures: list[str]) -> None:
    if not page_needs_human_verification(page):
        return
    print("Complete the site's human verification in the visible browser; crawling will continue automatically.", flush=True)
    failure_count = 0
    deadline = time.monotonic() + 300
    try:
        while page_needs_human_verification(page):
            while failure_count < len(challenge_failures):
                print(f"A verification request failed: {challenge_failures[failure_count]}", flush=True)
                failure_count += 1
            if time.monotonic() >= deadline:
                raise HumanVerificationError("Human verification did not clear within five minutes.")
            page.wait_for_timeout(1_000)
    except PlaywrightError as error:
        if page.is_closed():
            raise HumanVerificationError("The browser tab was closed before verification completed.") from error
        raise


def open_site(page: Page, config: CrawlerConfig, challenge_failures: list[str]) -> None:
    if not is_site_url(page.url, config):
        page.goto(config.site_url, wait_until="domcontentloaded", timeout=60_000)
    else:
        try:
            page.wait_for_load_state("domcontentloaded", timeout=30_000)
        except PlaywrightTimeoutError:
            pass
    wait_for_human_verification(page, challenge_failures)
    dismiss_blocking_popups(page)
    if find_search_box(page) is None:
        excerpt = page.locator("body").inner_text(timeout=15_000).strip().replace("\n", " ")[:300]
        raise RuntimeError(f"{config.brand} search was not found on {page.url!r} ({page.title()!r}). Page text: {excerpt!r}")


def product_part_number(url: str, config: CrawlerConfig) -> str:
    path = unquote(urlparse(url).path)
    match = re.search(r"[~_]([^/~_.]+(?:-[^/~_.]+)*)\.html?$", path, re.IGNORECASE)
    if match:
        return match.group(1)
    body_match = re.search(config.part_label_pattern, path, re.IGNORECASE)
    return body_match.group(1) if body_match else ""


def is_product_detail_page(url: str, config: CrawlerConfig) -> bool:
    parsed = urlparse(url)
    return bool(
        parsed.hostname == urlparse(config.site_url).hostname
        and re.match(r"^/parts/.+[~_][^/]+\.html?$", parsed.path, re.IGNORECASE)
    )


def detail_url_matches_oem(url: str, oem_number: object, config: CrawlerConfig) -> bool:
    actual = product_part_number(url, config)
    return bool(actual and part_numbers_match(actual, oem_number))


def find_matching_result(page: Page, oem_number: object, config: CrawlerConfig) -> Locator | None:
    expected = normalize_part_number(oem_number)
    if not expected:
        return None
    links = page.locator('a[href*="/parts/"]')
    for index in range(links.count()):
        link = links.nth(index)
        try:
            if not link.is_visible():
                continue
            label = link.inner_text().strip()
            href = link.get_attribute("href") or ""
        except PlaywrightError:
            continue
        if part_numbers_match(label, oem_number) or detail_url_matches_oem(urljoin(page.url, href), oem_number, config):
            return link
    return None


def is_matching_detail_page(page: Page, oem_number: object, config: CrawlerConfig) -> bool:
    if detail_url_matches_oem(page.url, oem_number, config):
        return True
    body_text = page.locator("body").inner_text(timeout=15_000)
    part_match = re.search(config.part_label_pattern, body_text, re.IGNORECASE)
    if part_match and part_numbers_match(part_match.group(1), oem_number):
        return True
    replaces = REPLACES_PATTERN.search(body_text)
    if replaces:
        replaced_numbers = re.findall(r"[A-Z0-9]+(?:-[A-Z0-9]+)*", replaces.group(1), re.IGNORECASE)
        return any(part_numbers_match(number, oem_number) for number in replaced_numbers)
    return False


def detail_status(body_text: str, oem_number: object, msrp: float | None, url: str, config: CrawlerConfig) -> str:
    if msrp is None:
        return "msrp_not_found"
    current_part = product_part_number(url, config)
    if not current_part:
        part_match = re.search(config.part_label_pattern, body_text, re.IGNORECASE)
        current_part = part_match.group(1) if part_match else ""
    if current_part and not part_numbers_match(current_part, oem_number):
        return f"success_superseded: {current_part}"
    return "success"


def wait_for_product_msrp(page: Page, timeout_seconds: float = 20) -> str:
    body = page.locator("body")
    deadline = time.monotonic() + timeout_seconds
    latest_text = ""
    while True:
        if body.count() and body.is_visible():
            latest_text = body.inner_text(timeout=2_000)
            if parse_msrp(latest_text) is not None:
                return latest_text
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return latest_text
        page.wait_for_timeout(min(500, max(1, int(remaining * 1_000))))


def crawl_msrp(page: Page, oem_number: object, challenge_failures: list[str], config: CrawlerConfig) -> tuple[float | None, str]:
    dismiss_blocking_popups(page)
    search_box = find_search_box(page)
    if search_box is None:
        if page_needs_human_verification(page):
            wait_for_human_verification(page, challenge_failures)
            dismiss_blocking_popups(page)
            search_box = find_search_box(page)
        if search_box is None:
            raise RuntimeError(f"{config.brand} search is unavailable on {page.url!r} ({page.title()!r}).")

    search_box.fill(str(oem_number).strip())
    starting_url = page.url
    search_box.press("Enter")
    try:
        page.wait_for_url(lambda current_url: current_url != starting_url, wait_until="domcontentloaded", timeout=12_000)
    except PlaywrightTimeoutError:
        pass
    wait_for_human_verification(page, challenge_failures)
    dismiss_blocking_popups(page)

    if not is_product_detail_page(page.url, config):
        result = find_matching_result(page, oem_number, config)
        if result is None:
            print(f"  Search did not open a {config.brand} detail page or matching result for {oem_number!s}; moving on.")
            return None, "not_found"
        href = result.get_attribute("href")
        if not href:
            return None, "not_found"
        try:
            page.goto(urljoin(page.url, href), wait_until="domcontentloaded", timeout=30_000)
        except PlaywrightTimeoutError:
            pass
        wait_for_human_verification(page, challenge_failures)
        dismiss_blocking_popups(page)

    if not is_product_detail_page(page.url, config) or not is_matching_detail_page(page, oem_number, config):
        print(f"  {config.brand} product page did not match OEM {oem_number!s}.")
        return None, "not_found"

    dismiss_blocking_popups(page)
    detail_text = wait_for_product_msrp(page)
    price = parse_msrp(detail_text)
    if price is None:
        print(f"  Detail page opened, but no MSRP label/value was found for {oem_number!s}.")
    else:
        print(f"  MSRP: ${price:,.2f}")
    return price, detail_status(detail_text, oem_number, price, page.url, config)


def save_results(output_path: Path, results: list[tuple[object, float | None, date, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(OUTPUT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for oem_number, msrp, crawled_on, status in results:
        sheet.append((oem_number, msrp, crawled_on, status))
        sheet.cell(sheet.max_row, 2).number_format = '$#,##0.00'
        sheet.cell(sheet.max_row, 3).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in (("A", 22), ("B", 14), ("C", 14), ("D", 32)):
        sheet.column_dimensions[column].width = width
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()


def load_results(output_path: Path) -> list[tuple[object, float | None, date, str]]:
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [str(value or "").strip().casefold() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        indexes = {name: header.index(name) for name in ("oem number", "msrp", "date", "status") if name in header}
        if "oem number" not in indexes:
            return []
        loaded = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            part = row[indexes["oem number"]] if indexes["oem number"] < len(row) else None
            if part is None or not str(part).strip():
                continue
            price = row[indexes["msrp"]] if "msrp" in indexes and indexes["msrp"] < len(row) else None
            crawled_on = row[indexes["date"]] if "date" in indexes and indexes["date"] < len(row) else date.today()
            status = str(row[indexes["status"]] or "").strip() if "status" in indexes and indexes["status"] < len(row) else ""
            loaded.append((part, price, crawled_on, status or ("success" if price is not None else "error: legacy row has no MSRP")))
        return loaded
    finally:
        workbook.close()


def prepare_crawl_rows(
    oem_values: list[object], previous: list[tuple[object, float | None, date, str]]
) -> tuple[list[tuple[object, float | None, date, str]], list[tuple[int, object]]]:
    previous_by_oem: dict[str, list[tuple[object, float | None, date, str]]] = {}
    for row in previous:
        previous_by_oem.setdefault(normalize_part_number(row[0]), []).append(row)
    results = []
    pending = []
    for row_index, part in enumerate(oem_values):
        prior_rows = previous_by_oem.get(normalize_part_number(part), [])
        prior = prior_rows.pop(0) if prior_rows else None
        if prior and prior[3].casefold().startswith("success"):
            results.append((part, prior[1], prior[2], prior[3]))
        else:
            results.append(prior if prior else (part, None, date.today(), "pending"))
            pending.append((row_index, part))
    return results, pending


def find_chrome(custom_path: Path | None = None) -> Path | None:
    if custom_path is not None:
        candidate = custom_path.expanduser().resolve()
        return candidate if candidate.is_file() else None
    candidates = [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "C:/Users/Default/AppData/Local")) / "Google/Chrome/Application/chrome.exe",
    ]
    return next((candidate for candidate in candidates if candidate.is_file()), None)


def cdp_is_ready(port: int) -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/json/version", timeout=1):
            return True
    except Exception:
        return False


def start_chrome(chrome_path: Path, profile_path: Path, config: CrawlerConfig) -> subprocess.Popen:
    profile_path.mkdir(parents=True, exist_ok=True)
    process = subprocess.Popen(
        [str(chrome_path), f"--remote-debugging-port={config.cdp_port}", f"--user-data-dir={profile_path}", "--no-first-run", "--no-default-browser-check", config.site_url],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(60):
        if cdp_is_ready(config.cdp_port):
            return process
        if process.poll() is not None:
            raise RuntimeError(f"Chrome exited during startup with code {process.returncode}.")
        time.sleep(0.5)
    process.terminate()
    raise RuntimeError(f"Chrome started but DevTools did not open on port {config.cdp_port}.")


def parse_args(description: str, argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument("--input", required=True, type=Path, help="Input .xlsx or .xlsm workbook.")
    parser.add_argument("--column", required=True, help="OEM column header, Excel letters, or 1-based number.")
    parser.add_argument("--output", type=Path, help="Output workbook (default: <input>_results.xlsx).")
    parser.add_argument("--sheet", help="Worksheet name (default: active worksheet).")
    parser.add_argument("--chrome-path", type=Path, help="Path to installed Google Chrome executable.")
    crawl_mode = parser.add_mutually_exclusive_group()
    crawl_mode.add_argument("--resume", dest="crawl_mode", action="store_const", const="resume", help="Resume (default).")
    crawl_mode.add_argument("--start-over", dest="crawl_mode", action="store_const", const="start-over", help="Recrawl all rows.")
    parser.set_defaults(crawl_mode="resume")
    parser.add_argument("--header-row", type=int, default=1, help="Header row number; use 0 for no header.")
    return parser.parse_args(argv)


def run(config: CrawlerConfig, script_dir: Path, description: str, argv: list[str] | None = None) -> int:
    args = parse_args(description, argv)
    input_path = args.input.expanduser().resolve()
    if not input_path.is_file() or input_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        print("Input must be an existing .xlsx or .xlsm workbook.", file=sys.stderr)
        return 2
    if args.header_row < 0:
        print("--header-row must be 0 or greater.", file=sys.stderr)
        return 2
    output_path = (args.output or input_path.with_name(f"{input_path.stem}_results.xlsx")).expanduser().resolve()
    if output_path == input_path:
        print("Input and output paths must be different.", file=sys.stderr)
        return 2
    try:
        oem_values = read_oem_values(input_path, args.sheet, args.column, args.header_row)
        previous = load_results(output_path) if args.crawl_mode == "resume" and output_path.is_file() else []
    except (OSError, ValueError, KeyError) as error:
        print(f"Could not prepare workbook: {error}", file=sys.stderr)
        return 2
    if not oem_values:
        print("No OEM values were found in the selected column.", file=sys.stderr)
        return 2
    results, pending = prepare_crawl_rows(oem_values, previous)
    if not pending:
        print("All input OEM numbers already have completed statuses. Use --start-over to crawl again.")
        return 0
    chrome_path = find_chrome(args.chrome_path)
    if chrome_path is None:
        print("Google Chrome was not found. Install Chrome or pass --chrome-path.", file=sys.stderr)
        return 2

    process = None
    browser = None
    keep_browser_open = False
    with sync_playwright() as playwright:
        process = start_chrome(chrome_path, script_dir / config.profile_name, config)
        try:
            browser = playwright.chromium.connect_over_cdp(f"http://127.0.0.1:{config.cdp_port}")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            pages = context.pages
            page = pages[0] if pages else context.new_page()
            for extra_page in pages[1:]:
                try:
                    extra_page.close()
                except PlaywrightError:
                    pass
            page.set_default_timeout(30_000)
            challenge_failures: list[str] = []

            def record_challenge_failure(request: Request) -> None:
                if "challenges.cloudflare.com" in request.url:
                    challenge_failures.append(f"{request.url} ({request.failure or 'request failed'})")

            page.on("requestfailed", record_challenge_failure)
            try:
                open_site(page, config, challenge_failures)
            except HumanVerificationError as error:
                keep_browser_open = True
                print(f"Crawl stopped: {error} Browser left open for verification.", file=sys.stderr)
                return 1
            for count, (row_index, part) in enumerate(pending, start=1):
                print(f"[{count}/{len(pending)}] Searching {part!s}")
                try:
                    price, status = crawl_msrp(page, part, challenge_failures, config)
                except HumanVerificationError as error:
                    keep_browser_open = True
                    print(f"Crawl stopped: {error} Browser left open for verification.", file=sys.stderr)
                    return 1
                except Exception as error:
                    print(f"  Crawl failed for {part!s}: {error}")
                    price, status = None, f"error: {type(error).__name__}: {error}"
                results[row_index] = (part, price, date.today(), status)
                save_results(output_path, results)
        finally:
            if browser is not None and not keep_browser_open:
                try:
                    browser.close()
                except Exception:
                    pass
            if process is not None and not keep_browser_open and process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    process.kill()
    print(f"Saved {len(results)} row(s) to {output_path}")
    return 0
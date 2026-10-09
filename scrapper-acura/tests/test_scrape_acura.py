import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

from scrape_acura import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    is_acura_site_url,
    is_product_detail_page,
    load_results,
    normalize_part_number,
    parse_args,
    parse_msrp,
    parse_product_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
    status_is_complete,
    wait_for_human_verification,
)


class AcuraCrawlerTests(unittest.TestCase):
    def test_parse_args_accepts_standard_options(self):
        with patch("sys.argv", ["scrape_acura.py", "--input", "parts.xlsx", "--column", "A"]):
            args = parse_args()
        self.assertEqual(args.input, Path("parts.xlsx"))
        self.assertEqual(args.column, "A")

    def test_normalize_oem(self):
        self.assertEqual(normalize_part_number("17220-5J2-A00"), "172205J2A00")

    def test_parse_msrp_reads_msrp_not_discounted_price(self):
        self.assertEqual(parse_msrp("$27.39 MSRP: $38.20\nYou Save: $10.81"), 38.20)

    def test_parse_product_msrp_from_acura_details(self):
        text = """Acura 17220-5J2-A00 Engine Air Filter Cleaner Element
Manufacturer Part Number
17220-5J2-A00
$27.39 MSRP: $38.20
Currently Unavailable"""
        self.assertEqual(parse_product_msrp(text), 38.20)

    def test_live_oem_60260_tgv_a00zz_msrp_is_extracted(self):
        live_price_text = """$338.08 MSRP: $482.83
You Save: $144.75 (30%)
Currently Unavailable
Due to manufacturer supply shortage, this item isn't currently available for ordering."""
        self.assertEqual(parse_msrp(live_price_text), 482.83)

    def test_parse_product_msrp_requires_acura_part_reference(self):
        self.assertIsNone(parse_product_msrp("MSRP: $38.20"))

    def test_detail_url_matches_acura_oem_route(self):
        url = "https://www.acurapartswarehouse.com/oem/acura~element~assy~air~17220-5j2-a00.html"
        self.assertTrue(detail_url_matches_oem(url, "17220-5J2-A00"))
        self.assertFalse(detail_url_matches_oem(url, "17220-5J2-A01"))
        self.assertFalse(detail_url_matches_oem("https://example.com/oem/item-17220-5j2-a00.html", "17220-5J2-A00"))

    def test_site_url_recognizes_home_and_product_pages(self):
        self.assertTrue(is_acura_site_url("https://www.acurapartswarehouse.com/"))
        self.assertTrue(is_acura_site_url("https://www.acurapartswarehouse.com/oem/acura~part~17220-5j2-a00.html"))
        self.assertFalse(is_acura_site_url("about:blank"))

    def test_direct_product_url_is_enough_to_identify_detail_page(self):
        page = Mock()
        page.url = "https://www.acurapartswarehouse.com/oem/acura~fender~assy~l~fr~60260-tgv-a00zz.html"
        self.assertTrue(is_product_detail_page(page))

    def test_status_identifies_superseded_manufacturer_number(self):
        self.assertEqual(
            detail_status("Manufacturer Part Number\n17220-5J2-A01", "17220-5J2-A00", 38.20),
            "success_superseded: 17220-5J2-A01",
        )
        self.assertEqual(detail_status("Acura part\nMSRP: $38.20", "17220-5J2-A00", 38.20), "success")

    def test_only_successful_statuses_are_complete(self):
        self.assertTrue(status_is_complete("success"))
        self.assertTrue(status_is_complete("success_superseded: 17220-5J2-A01"))
        self.assertFalse(status_is_complete("not_found"))
        self.assertFalse(status_is_complete("msrp_not_found"))

    def test_missing_suggestion_submits_full_search_once(self):
        page = Mock()
        page.url = "https://www.acurapartswarehouse.com/"
        search_box = Mock()
        with (
            patch("scrape_acura.dismiss_blocking_popups"),
            patch("scrape_acura.find_search_box", return_value=search_box),
            patch("scrape_acura.find_search_button", return_value=None),
            patch("scrape_acura.wait_for_matching_result", return_value=None),
            patch("scrape_acura.wait_for_human_verification"),
            patch("scrape_acura.is_product_detail_page", return_value=False),
        ):
            result = crawl_msrp(page, "17220-5J2-A00", [])
        self.assertEqual(result, (None, "not_found"))
        search_box.fill.assert_called_once_with("17220-5J2-A00")
        search_box.press.assert_called_once_with("Enter")

    def test_search_uses_acura_submit_button(self):
        page = Mock()
        page.url = "https://www.acurapartswarehouse.com/"
        search_box = Mock()
        search_button = Mock()
        with (
            patch("scrape_acura.dismiss_blocking_popups"),
            patch("scrape_acura.find_search_box", return_value=search_box),
            patch("scrape_acura.find_search_button", return_value=search_button),
            patch("scrape_acura.wait_for_human_verification"),
            patch("scrape_acura.is_product_detail_page", return_value=True),
            patch("scrape_acura.is_matching_detail_page", return_value=True),
            patch("scrape_acura.wait_for_product_msrp", return_value="$27.39 MSRP: $38.20"),
        ):
            result = crawl_msrp(page, "17220-5J2-A00", [])
        self.assertEqual(result, (38.20, "success"))
        search_button.click.assert_called_once_with(timeout=10_000)
        search_box.press.assert_not_called()

    def test_direct_search_navigation_extracts_live_oem_msrp(self):
        page = Mock()
        page.url = "https://www.acurapartswarehouse.com/"
        search_box = Mock()
        search_button = Mock()
        product_url = (
            "https://www.acurapartswarehouse.com/oem/"
            "acura~fender~assy~l~fr~60260-tgv-a00zz.html"
        )
        search_button.click.side_effect = lambda **_: setattr(page, "url", product_url)
        price_text = "$338.08 MSRP: $482.83\\nYou Save: $144.75 (30%)"
        with (
            patch("scrape_acura.dismiss_blocking_popups") as dismiss_popups,
            patch("scrape_acura.find_search_box", return_value=search_box),
            patch("scrape_acura.find_search_button", return_value=search_button),
            patch("scrape_acura.wait_for_human_verification"),
            patch("scrape_acura.wait_for_product_msrp", return_value=price_text),
        ):
            result = crawl_msrp(page, "60260TGVA00ZZ", [])
        self.assertEqual(result, (482.83, "success"))
        page.wait_for_url.assert_called_once()
        search_button.click.assert_called_once_with(timeout=10_000)
        self.assertGreaterEqual(dismiss_popups.call_count, 3)

    def test_human_verification_waits_without_reloading(self):
        page = Mock()
        failures = ["challenge request failed"]
        with patch("scrape_acura.page_needs_human_verification", side_effect=[True, True, False]):
            wait_for_human_verification(page, failures)
        page.reload.assert_not_called()
        page.wait_for_timeout.assert_called_once_with(1_000)

    def test_resolve_column_supports_header_and_excel_letter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("Description", "OEM Number"))
        sheet.append(("Air filter", "17220-5J2-A00"))
        self.assertEqual(resolve_column(sheet, "OEM Number", 1), (2, 2))
        self.assertEqual(resolve_column(sheet, "B", 1), (2, 2))
        workbook.close()

    def test_read_oem_values(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "parts.xlsx"
            workbook = Workbook()
            workbook.active.append(("Description", "OEM Number"))
            workbook.active.append(("Air filter", "17220-5J2-A00"))
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_oem_values(path, None, "OEM Number", 1), ["17220-5J2-A00"])

    def test_duplicate_oems_remain_separate_and_resume_by_occurrence(self):
        saved = [
            ("17220-5J2-A00", 38.20, date(2026, 10, 9), "success"),
            ("17220-5J2-A00", None, date(2026, 10, 9), "not_found"),
        ]
        results, pending = prepare_crawl_rows(["17220-5J2-A00", "17220-5J2-A00"], saved)
        self.assertEqual(results[0][1:], saved[0][1:])
        self.assertEqual(pending, [(1, "17220-5J2-A00")])

    def test_resume_recrawls_success_status_without_msrp(self):
        saved = [("60260TGVA00ZZ", None, date(2026, 10, 9), "success")]
        results, pending = prepare_crawl_rows(["60260TGVA00ZZ"], saved)
        self.assertEqual(results[0][1], None)
        self.assertEqual(pending, [(0, "60260TGVA00ZZ")])

    def test_save_and_load_results_round_trip(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            rows = [("17220-5J2-A00", 38.20, date(2026, 10, 9), "success")]
            save_results(output, rows)
            self.assertEqual(load_results(output)[0][:2], ("17220-5J2-A00", 38.20))
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(tuple(cell.value for cell in workbook.active[1]), ("OEM Number", "MSRP", "Date", "Status"))
            workbook.close()


if __name__ == "__main__":
    unittest.main()

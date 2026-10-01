import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from scrape_vw import (
    load_results,
    detail_status,
    detail_url_matches_oem,
    normalize_part_number,
    parse_msrp,
    read_oem_values,
    resolve_column,
    save_results,
    status_is_complete,
)


class CrawlerHelpersTests(unittest.TestCase):
    def test_normalize_part_number_ignores_punctuation_and_case(self):
        self.assertEqual(normalize_part_number("1K0-123-456.A"), "1K0123456A")

    def test_detail_url_matches_queried_superseded_oem(self):
        self.assertTrue(
            detail_url_matches_oem(
                "https://parts.vw.com/p/48644238/1K9853651AZLL.html",
                "1K9853651AZLL",
            )
        )
        self.assertFalse(
            detail_url_matches_oem(
                "https://parts.vw.com/p/48644238/1K9853651EZLL.html",
                "1K9853651AZLL",
            )
        )

    def test_superseded_detail_status_discloses_current_part_number(self):
        self.assertEqual(
            detail_status("Part Number: 1K9853651EZLL", "1K9853651AZLL", 429.04),
            "success_superseded: 1K9853651EZLL",
        )

    def test_only_successful_statuses_are_complete(self):
        self.assertTrue(status_is_complete("success"))
        self.assertTrue(status_is_complete("success_superseded: 1K9853651EZLL"))
        self.assertFalse(status_is_complete("not_found"))
        self.assertFalse(status_is_complete("msrp_not_found"))
        self.assertFalse(status_is_complete("error: TimeoutError"))

    def test_parse_msrp_reads_currency_and_thousands_separator(self):
        self.assertEqual(parse_msrp("Part information MSRP: $1,234.50"), 1234.5)
        self.assertEqual(parse_msrp("MSRP\n$ 553.59"), 553.59)

    def test_parse_msrp_requires_an_msrp_label(self):
        self.assertIsNone(parse_msrp("Price: $123.45"))

    def test_resolve_column_supports_header_and_excel_letter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("Description", "OEM Number"))
        sheet.append(("Widget", "1K0123456A"))
        self.assertEqual(resolve_column(sheet, "OEM Number", 1), (2, 2))
        self.assertEqual(resolve_column(sheet, "B", 1), (2, 2))

    def test_read_oem_values_handles_headerless_sheet(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "parts.xlsx"
            workbook = Workbook()
            workbook.active.append(("1K0123456A",))
            workbook.active.append(("5Q0123456B",))
            workbook.save(path)
            self.assertEqual(read_oem_values(path, None, "A", 0), ["1K0123456A", "5Q0123456B"])

    def test_save_results_creates_status_column_and_round_trips_rows(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            results = [
                ("1K0123456A", 12.5, date(2026, 10, 1), "success"),
                ("5Q0123456B", None, date(2026, 10, 1), "error: TimeoutError"),
            ]
            save_results(output, results)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                tuple(cell.value for cell in workbook.active[1]),
                ("OEM Number", "MSRP", "Date", "Status"),
            )
            row = tuple(cell.value for cell in workbook.active[2])
            self.assertEqual(row[:2], ("1K0123456A", 12.5))
            self.assertEqual(row[2].date(), date(2026, 10, 1))
            self.assertEqual(row[3], "success")
            workbook.close()

            loaded = load_results(output)
            self.assertEqual(loaded[0][0:2], ("1K0123456A", 12.5))
            self.assertEqual(loaded[0][3], "success")
            self.assertEqual(loaded[1][3], "error: TimeoutError")
            self.assertTrue(status_is_complete(loaded[0][3]))
            self.assertFalse(status_is_complete(loaded[1][3]))

    def test_legacy_rows_without_status_retry_when_msrp_is_empty(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "legacy.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("OEM Number", "MSRP", "Date"))
            sheet.append(("1K0123456A", None, date(2026, 10, 1)))
            sheet.append(("5Q0123456B", 19.5, date(2026, 10, 1)))
            workbook.save(output)
            loaded = load_results(output)
            self.assertEqual(loaded[0][3], "error: legacy row has no status or MSRP")
            self.assertEqual(loaded[1][3], "success")


if __name__ == "__main__":
    unittest.main()
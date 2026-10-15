import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scrape_porsche import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    find_matching_result,
    is_matching_detail_page,
    load_results,
    normalize_part_number,
    parse_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
)


class PorscheCrawlerTests(unittest.TestCase):
    def test_normalize_part_number(self):
        self.assertEqual(normalize_part_number("7PP-035-223-AA"), "7PP035223AA")

    def test_parse_msrp_from_product_detail(self):
        self.assertEqual(parse_msrp("Price $1,506.41\nMSRP $1,909.39"), 1909.39)
        self.assertIsNone(parse_msrp("Price: $24.50"))

    def test_detail_url_matches_compressed_part_number(self):
        self.assertTrue(
            detail_url_matches_oem(
                "https://porsche.oempartsonline.com/oem-parts/porsche-radio-amplifier-7pp035223aa",
                "7PP-035-223-AA",
            )
        )
        self.assertFalse(
            detail_url_matches_oem(
                "https://example.com/oem-parts/porsche-radio-amplifier-7pp035223aa",
                "7PP-035-223-AA",
            )
        )

    def test_result_match_prefers_exact_displayed_part_number(self):
        partial = Mock()
        partial.is_visible.return_value = True
        partial.inner_text.return_value = "7PP-035-223-AAB"
        partial.get_attribute.return_value = "/oem-parts/porsche-radio-amplifier-7pp035223aab"
        exact = Mock()
        exact.is_visible.return_value = True
        exact.inner_text.return_value = "7PP-035-223-AA"
        exact.get_attribute.return_value = "/oem-parts/porsche-radio-amplifier-7pp035223aa"
        links = Mock()
        links.count.return_value = 2
        links.nth.side_effect = lambda index: (partial, exact)[index]
        page = Mock()
        page.locator.return_value = links

        self.assertIs(find_matching_result(page, "7PP-035-223-AA"), exact)

    def test_result_match_does_not_accept_a_prefix_only(self):
        link = Mock()
        link.is_visible.return_value = True
        link.inner_text.return_value = "7PP-035-223-AAB"
        link.get_attribute.return_value = "/oem-parts/porsche-radio-amplifier-7pp035223aab"
        links = Mock()
        links.count.return_value = 1
        links.nth.return_value = link
        page = Mock()
        page.url = "about:blank"
        page.locator.return_value = links

        self.assertIsNone(find_matching_result(page, "7PP-035-223-AA"))

    def test_detail_page_matches_mpn(self):
        page = Mock()
        page.url = "https://porsche.oempartsonline.com/oem-parts/porsche-radio-amplifier"
        page.locator.return_value.inner_text.return_value = "Radio Amplifier - Porsche\nMPN 7pp035223aa\nMSRP $1,909.39"

        self.assertTrue(is_matching_detail_page(page, "7PP-035-223-AA"))

    def test_crawl_opens_matching_product_and_reads_msrp(self):
        page = Mock()
        page.url = "https://porsche.oempartsonline.com/"
        page.locator.return_value.inner_text.return_value = "Radio Amplifier - Porsche (7PP-035-223-AA)\nMPN 7pp035223aa\nMSRP $1,909.39"
        search_box = Mock()
        matching_result = Mock()
        with (
            patch("scrape_porsche.dismiss_blocking_popups"),
            patch("scrape_porsche.find_search_box", return_value=search_box),
            patch("scrape_porsche.wait_for_human_verification"),
            patch("scrape_porsche.is_matching_detail_page", side_effect=[False, True]),
            patch("scrape_porsche.wait_for_matching_result", return_value=matching_result),
        ):
            result = crawl_msrp(page, "7PP-035-223-AA", [])

        self.assertEqual(result, (1909.39, "success"))
        search_box.fill.assert_called_once_with("7PP-035-223-AA")
        matching_result.click.assert_called_once_with(timeout=10_000)

    def test_detail_status_reports_replacement_mpn(self):
        self.assertEqual(
            detail_status("Radio Amplifier\nMPN 7PP035223AB", "7PP035223AA", 1909.39),
            "success_superseded: 7PP035223AB",
        )

    def test_excel_column_and_result_workbook_round_trip(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            workbook = Workbook()
            workbook.active.append(("Description", "OEM Number"))
            workbook.active.append(("Radio amplifier", "7PP-035-223-AA"))
            workbook.save(output)
            workbook.close()
            self.assertEqual(read_oem_values(output, None, "OEM Number", 1), ["7PP-035-223-AA"])
            source_book = load_workbook(output)
            self.assertEqual(resolve_column(source_book.active, "B", 1), (2, 2))
            source_book.close()

            rows = [("7PP-035-223-AA", 1909.39, date(2026, 10, 14), "success")]
            result_path = output.with_name("saved.xlsx")
            save_results(result_path, rows)
            loaded = load_results(result_path)
            self.assertEqual(loaded[0][:2], rows[0][:2])
            saved_book = load_workbook(result_path, data_only=True)
            self.assertEqual(tuple(cell.value for cell in saved_book.active[1]), ("OEM Number", "MSRP", "Date", "Status"))
            saved_book.close()

    def test_resume_skips_successful_occurrences_only(self):
        saved = [("7PP-035-223-AA", 1909.39, date(2026, 10, 14), "success")]
        results, pending = prepare_crawl_rows(["7PP-035-223-AA", "7PP-035-223-AA"], saved)
        self.assertEqual(results[0][3], "success")
        self.assertEqual(pending, [(1, "7PP-035-223-AA")])


if __name__ == "__main__":
    unittest.main()
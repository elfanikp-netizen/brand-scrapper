import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scrape_bmw import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    find_matching_result,
    is_matching_detail_page,
    load_results,
    normalize_part_number,
    parse_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
)


class BMWCrawlerTests(unittest.TestCase):
    def test_normalize_part_number(self):
        self.assertEqual(normalize_part_number("11-42 7 566 327"), "11427566327")

    def test_parse_msrp_from_detail_page(self):
        self.assertEqual(parse_msrp("Price $18.25\nMSRP: $24.50"), 24.5)
        self.assertEqual(parse_msrp("MSRP $1,240.00"), 1240.0)
        self.assertIsNone(parse_msrp("Price: $24.50"))

    def test_detail_url_matches_part_number(self):
        self.assertTrue(detail_url_matches_oem("https://www.getbmwparts.com/oem-parts/bmw-11427566327", "11-42-7-566-327"))
        self.assertFalse(detail_url_matches_oem("https://example.com/oem-parts/bmw-11427566327", "11-42-7-566-327"))

    def test_result_match_requires_exact_part_number_and_prefers_displayed_number(self):
        partial = Mock()
        partial.is_visible.return_value = True
        partial.inner_text.return_value = "Part 11-42-7-566-3271 filter"
        partial.get_attribute.return_value = "/oem-parts/bmw-114275663271"
        exact = Mock()
        exact.is_visible.return_value = True
        exact.inner_text.return_value = "BMW 11-42-7-566-327 Oil Filter"
        exact.get_attribute.return_value = "/oem-parts/bmw-11427566327"
        locator = Mock()
        locator.count.return_value = 2
        locator.nth.side_effect = lambda index: (partial, exact)[index]
        page = Mock()
        page.locator.return_value = locator

        self.assertIs(find_matching_result(page, "11-42-7-566-327"), exact)

    def test_result_match_does_not_accept_a_prefix_only(self):
        link = Mock()
        link.is_visible.return_value = True
        link.inner_text.return_value = "11-42-7-566-3271"
        link.get_attribute.return_value = "/oem-parts/114275663271"
        locator = Mock()
        locator.count.return_value = 1
        locator.nth.return_value = link
        page = Mock()
        page.url = "about:blank"
        page.locator.return_value = locator

        self.assertIsNone(find_matching_result(page, "11-42-7-566-327"))

    def test_single_result_on_exact_search_page_can_be_a_replacement_part(self):
        replacement = Mock()
        replacement.is_visible.return_value = True
        replacement.inner_text.return_value = "41-00-9-633-544"
        replacement.get_attribute.return_value = "/oem-parts/bmw-fender-41009633544"
        links = Mock()
        links.count.return_value = 1
        links.nth.return_value = replacement
        page = Mock()
        page.url = "https://www.getbmwparts.com/search?search_str=41007496138"
        page.locator.return_value = links

        self.assertIs(find_matching_result(page, "41007496138"), replacement)

    def test_multiple_close_results_choose_smallest_part_number_distance(self):
        closest = Mock()
        closest.is_visible.return_value = True
        closest.inner_text.return_value = "63-21-7-888-850"
        closest.get_attribute.return_value = "/oem-parts/bmw-back-up-light-assembly-63217888850"
        next_closest = Mock()
        next_closest.is_visible.return_value = True
        next_closest.inner_text.return_value = "63-21-7-888-852"
        next_closest.get_attribute.return_value = "/oem-parts/bmw-rear-light-63217888852"
        links = Mock()
        links.count.return_value = 2
        links.nth.side_effect = lambda index: (closest, next_closest)[index]
        page = Mock()
        page.url = "https://www.getbmwparts.com/search?search_str=6321788885"
        page.locator.return_value = links

        self.assertIs(find_matching_result(page, "6321788885"), closest)

    def test_detail_page_replaces_field_confirms_searched_oem(self):
        page = Mock()
        page.url = "https://www.getbmwparts.com/oem-parts/bmw-fender-41009633544"
        page.locator.return_value.inner_text.return_value = (
            "Genuine BMW Fender - Part Number 41-00-9-633-544\n"
            "MSRP: $595.37\nReplaces: 41007496138"
        )

        self.assertTrue(is_matching_detail_page(page, "41007496138"))

    def test_crawl_clicks_matching_result_before_reading_detail_msrp(self):
        page = Mock()
        page.url = "https://www.getbmwparts.com/"
        page.locator.return_value.inner_text.return_value = "Part Number: 11-42-7-566-327\nMSRP: $24.50"
        search_box = Mock()
        matching_result = Mock()
        with (
            patch("scrape_bmw.find_search_box", return_value=search_box),
            patch("scrape_bmw.wait_for_human_verification"),
            patch("scrape_bmw.is_matching_detail_page", side_effect=[False, True]),
            patch("scrape_bmw.wait_for_matching_result", return_value=matching_result),
        ):
            result = crawl_msrp(page, "11-42-7-566-327", [])

        self.assertEqual(result, (24.5, "success"))
        search_box.fill.assert_called_once_with("11-42-7-566-327")
        matching_result.click.assert_called_once_with(timeout=10_000)

    def test_crawl_accepts_selected_candidate_detail_without_replaces_field(self):
        page = Mock()
        page.url = "https://www.getbmwparts.com/search?search_str=6321788885"
        search_box = Mock()
        matching_result = Mock()
        matching_result.get_attribute.return_value = "/oem-parts/bmw-back-up-light-assembly-63217888850"
        matching_result.click.side_effect = lambda **_: setattr(
            page, "url", "https://www.getbmwparts.com/oem-parts/bmw-back-up-light-assembly-63217888850"
        )
        page.locator.return_value.inner_text.return_value = (
            "Genuine BMW Back Up Light Assembly - Part Number 63-21-7-888-850\nMSRP: $304.22"
        )
        with (
            patch("scrape_bmw.find_search_box", return_value=search_box),
            patch("scrape_bmw.wait_for_human_verification"),
            patch("scrape_bmw.is_matching_detail_page", return_value=False),
            patch("scrape_bmw.wait_for_matching_result", return_value=matching_result),
        ):
            result = crawl_msrp(page, "6321788885", [])

        self.assertEqual(result, (304.22, "success_superseded: 63-21-7-888-850"))

    def test_detail_status_reports_superseded_part(self):
        self.assertEqual(
            detail_status("Genuine BMW Fender - Part Number 41-00-9-633-544", "41007496138", 595.37),
            "success_superseded: 41-00-9-633-544",
        )

    def test_excel_column_and_result_workbook_round_trip(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            workbook = Workbook()
            workbook.active.append(("Description", "OEM Number"))
            workbook.active.append(("Oil filter", "11427566327"))
            workbook.save(output)
            workbook.close()
            self.assertEqual(read_oem_values(output, None, "OEM Number", 1), ["11427566327"])
            source_book = load_workbook(output)
            self.assertEqual(resolve_column(source_book.active, "B", 1), (2, 2))
            source_book.close()

            rows = [("11427566327", 24.5, date(2026, 10, 12), "success")]
            save_results(output.with_name("saved.xlsx"), rows)
            loaded = load_results(output.with_name("saved.xlsx"))
            self.assertEqual(loaded[0][:2], rows[0][:2])
            saved_book = load_workbook(output.with_name("saved.xlsx"), data_only=True)
            self.assertEqual(tuple(cell.value for cell in saved_book.active[1]), ("OEM Number", "MSRP", "Date", "Status"))
            saved_book.close()

    def test_resume_skips_successful_occurrences_only(self):
        saved = [("11427566327", 24.5, date(2026, 10, 12), "success")]
        results, pending = prepare_crawl_rows(["11427566327", "11427566327"], saved)
        self.assertEqual(results[0][3], "success")
        self.assertEqual(pending, [(1, "11427566327")])


if __name__ == "__main__":
    unittest.main()
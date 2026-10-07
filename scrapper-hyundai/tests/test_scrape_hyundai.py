import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

from scrape_hyundai import (
    cdp_page_is_hyundai,
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    is_hyundai_site_url,
    load_results,
    normalize_part_number,
    parse_args,
    parse_msrp,
    parse_product_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
    should_start_chrome,
    status_is_complete,
    wait_for_human_verification,
)


class HyundaiCrawlerTests(unittest.TestCase):
    def test_parse_args_supports_basic_crawl_options(self):
        with patch("sys.argv", ["scrape_hyundai.py", "--input", "parts.xlsx", "--column", "A"]):
            args = parse_args()
        self.assertEqual(args.input, Path("parts.xlsx"))
        self.assertEqual(args.column, "A")

    def test_missing_suggestion_submits_search_once(self):
        page = Mock()
        search_box = Mock()
        page.title.return_value = "Hyundai Parts"
        page.url = "https://www.hyundaipartsdeal.com/"
        page.locator.return_value.inner_text.return_value = "Hyundai Parts"
        with (
            patch("scrape_hyundai.dismiss_blocking_popups"),
            patch("scrape_hyundai.find_search_box", return_value=search_box),
            patch("scrape_hyundai.wait_for_matching_result", return_value=None),
            patch("scrape_hyundai.page_needs_human_verification", return_value=False),
        ):
            result = crawl_msrp(page, "28113-2W100", [])
        self.assertEqual(result, (None, "not_found"))
        search_box.fill.assert_called_once_with("28113-2W100")
        self.assertEqual(search_box.press.call_count, 1)

    def test_search_uses_enter_then_clicks_button_when_needed(self):
        page = Mock()
        search_box = Mock()
        search_button = Mock()
        page.title.return_value = "Hyundai Parts"
        page.url = "https://www.hyundaipartsdeal.com/"
        page.locator.return_value.inner_text.return_value = "Hyundai Parts"
        page.wait_for_url.side_effect = PlaywrightTimeoutError("no navigation")
        with (
            patch("scrape_hyundai.dismiss_blocking_popups"),
            patch("scrape_hyundai.find_search_box", return_value=search_box),
            patch("scrape_hyundai.find_search_button", return_value=search_button),
            patch("scrape_hyundai.is_product_detail_page", return_value=True),
            patch("scrape_hyundai.wait_for_product_detail", return_value=True),
            patch("scrape_hyundai.is_matching_detail_page", return_value=True),
            patch("scrape_hyundai.wait_for_product_msrp", return_value="Hyundai 28113-2W100\nMSRP: $30.06"),
            patch("scrape_hyundai.parse_product_msrp", return_value=30.06),
            patch("scrape_hyundai.detail_status", return_value="success"),
        ):
            result = crawl_msrp(page, "28113-2W100", [])
        self.assertEqual(result, (30.06, "success"))
        search_box.fill.assert_called_once_with("28113-2W100")
        search_box.press.assert_called_once_with("Enter")
        search_button.click.assert_called_once_with(timeout=10_000)

    def test_verification_waits_for_manual_completion_after_request_failure(self):
        page = Mock()
        failures = ["challenge request failed: net::ERR_NAME_NOT_RESOLVED"]
        with patch(
            "scrape_hyundai.page_needs_human_verification",
            side_effect=[True, True, False],
        ):
            wait_for_human_verification(page, failures)
        page.reload.assert_not_called()
        self.assertEqual(page.wait_for_timeout.call_count, 1)

    def test_normalize_part_number(self):
        self.assertEqual(normalize_part_number("28113-2W100"), "281132W100")

    def test_parse_msrp_with_currency_and_thousands_separator(self):
        self.assertEqual(parse_msrp("Hyundai 28113-2W100\nMSRP: $1,030.06"), 1030.06)

    def test_parse_product_msrp_from_live_hyundai_detail_text(self):
        text = """Hyundai 28113-2W100 Air Cleaner Filter
Part Description
Filter-Air Cleaner
$23.26 MSRP: $30.06
You Save: $6.80 (23%)
Related Parts
Hyundai 28110-B8100
MSRP $999.99"""
        self.assertEqual(parse_product_msrp(text), 30.06)

    def test_parse_product_msrp_falls_back_to_list_price_without_discount(self):
        self.assertEqual(parse_product_msrp("Hyundai 28113-2W100 Air Cleaner\nMSRP $30.06"), 30.06)

    def test_parse_product_msrp_requires_hyundai_product_heading(self):
        self.assertIsNone(parse_product_msrp("MSRP $30.06"))

    def test_detail_url_matches_hyundai_genuine_part_url(self):
        self.assertTrue(
            detail_url_matches_oem(
                "https://www.hyundaipartsdeal.com/genuine/hyundai-filter-air-cleaner~28113-2w100.html",
                "28113-2W100",
            )
        )
        self.assertFalse(
            detail_url_matches_oem(
                "https://www.hyundaipartsdeal.com/genuine/hyundai-filter-air-cleaner~28113-2w100.html",
                "28113-2W101",
            )
        )

    def test_site_url_recognizes_home_and_product_pages(self):
        self.assertTrue(is_hyundai_site_url("https://www.hyundaipartsdeal.com/"))
        self.assertTrue(
            is_hyundai_site_url(
                "https://www.hyundaipartsdeal.com/genuine/hyundai-filter-air-cleaner~28113-2w100.html"
            )
        )
        self.assertFalse(is_hyundai_site_url("about:blank"))

    def test_cdp_page_rejects_stale_non_hyundai_session(self):
        response = Mock()
        response.read.return_value = b'[{"url": "https://www.kiapartsnow.com/"}]'
        with patch("scrape_hyundai.urllib.request.urlopen", return_value=response):
            self.assertFalse(cdp_page_is_hyundai())

    def test_should_start_chrome_reuses_existing_hyundai_browser(self):
        with (
            patch("scrape_hyundai.cdp_is_ready", return_value=True),
            patch("scrape_hyundai.cdp_page_is_hyundai", return_value=True),
        ):
            self.assertFalse(should_start_chrome())

    def test_should_start_chrome_starts_new_browser_for_stale_or_non_hyundai_session(self):
        with (
            patch("scrape_hyundai.cdp_is_ready", return_value=True),
            patch("scrape_hyundai.cdp_page_is_hyundai", return_value=False),
        ):
            self.assertTrue(should_start_chrome())

    def test_status_reports_superseded_hyundai_part(self):
        self.assertEqual(
            detail_status("Hyundai 28113-2W101 Air Cleaner Filter", "28113-2W100", 30.06),
            "success_superseded: 28113-2W101",
        )
        self.assertEqual(detail_status("Hyundai Parts\nMSRP $30.06", "28113-2W100", 30.06), "success")
        self.assertEqual(detail_status("Clearance\nMSRP $30.06", "28113-2W100", 30.06), "success")

    def test_only_successful_statuses_are_complete(self):
        self.assertTrue(status_is_complete("success"))
        self.assertTrue(status_is_complete("success_superseded: 28113-2W101"))
        self.assertFalse(status_is_complete("not_found"))
        self.assertFalse(status_is_complete("msrp_not_found"))
        self.assertFalse(status_is_complete("error: TimeoutError"))

    def test_resolve_column_supports_header_and_excel_letter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("Description", "OEM Number"))
        sheet.append(("Air filter", "28113-2W100"))
        self.assertEqual(resolve_column(sheet, "OEM Number", 1), (2, 2))
        self.assertEqual(resolve_column(sheet, "B", 1), (2, 2))
        workbook.close()

    def test_read_oem_values_from_workbook(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "parts.xlsx"
            workbook = Workbook()
            workbook.active.append(("Description", "OEM Number"))
            workbook.active.append(("Air filter", "28113-2W100"))
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_oem_values(path, None, "OEM Number", 1), ["28113-2W100"])

    def test_duplicate_oems_are_queued_as_separate_rows(self):
        results, pending = prepare_crawl_rows(["28113-2W100", "28113-2W100"], [])
        self.assertEqual(len(results), 2)
        self.assertEqual(pending, [(0, "28113-2W100"), (1, "28113-2W100")])

    def test_resume_matches_duplicate_results_by_occurrence(self):
        saved = [
            ("28113-2W100", 30.06, date(2026, 10, 5), "success"),
            ("28113-2W100", None, date(2026, 10, 5), "not_found"),
        ]
        results, pending = prepare_crawl_rows(["28113-2W100", "28113-2W100"], saved)
        self.assertEqual(results[0][1:], saved[0][1:])
        self.assertEqual(pending, [(1, "28113-2W100")])

    def test_save_and_load_results_round_trip(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            rows = [("28113-2W100", 30.06, date(2026, 10, 5), "success")]
            save_results(output, rows)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                tuple(cell.value for cell in workbook.active[1]),
                ("OEM Number", "MSRP", "Date", "Status"),
            )
            workbook.close()
            self.assertEqual(load_results(output)[0][0:2], ("28113-2W100", 30.06))


if __name__ == "__main__":
    unittest.main()
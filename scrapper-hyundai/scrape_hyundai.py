from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import time
import urllib.request
from datetime import date
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string
from playwright.sync_api import (
    Error as PlaywrightError,
    Locator,
    Page,
    Request,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)


SITE_URL = "https://www.hyundaipartsdeal.com/"
SEARCH_INPUT_SELECTOR = 'input[placeholder="Search by Part Number, Part Name, or Vehicle"]'
SEARCH_BUTTON_SELECTOR = 'button[aria-label="Search by Part Number, Part Name, or Vehicle"]'
CDP_PORT = 9236
OUTPUT_HEADERS = ("OEM Number", "MSRP", "Date", "Status")
MSRP_PATTERN = re.compile(
    r"\bMSRP\b\s*(?:\([^)]*\))?\s*:?\s*\$?\s*([0-9][0-9,]*(?:\.\d{1,2})?)",
    re.IGNORECASE,
)


class HumanVerificationError(RuntimeError):
    pass


def normalize_part_number(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def parse_msrp(text: str) -> float | None:
    match = MSRP_PATTERN.search(text)
    if match is None:
        return None
    return float(match.group(1).replace(",", ""))


def parse_product_msrp(text: str) -> float | None:
    product_heading = re.search(r"(?i)Hyundai\s+#?[A-Z0-9-]+", text)
    if product_heading is None:
        return None
    product_text = text[product_heading.start() : product_heading.start() + 800]
    msrp_match = MSRP_PATTERN.search(product_text)
    if msrp_match is None:
        return None
    return float(msrp_match.group(1).replace(",", ""))


def resolve_column(sheet, column: str, header_row: int) -> tuple[int, int]:
    """Return the 1-based input column and first data row."""
    if header_row > 0:
        for index, cell in enumerate(sheet[header_row], start=1):
            if str(cell.value or "").strip().casefold() == column.strip().casefold():
                return index, header_row + 1

    if column.isdecimal():
        column_index = int(column)
        if column_index < 1:
            raise ValueError("A numeric --column value must be 1 or greater.")
    else:
        try:
            column_index = column_index_from_string(column.strip().upper())
        except ValueError as error:
            raise ValueError(
                f"Column {column!r} was not found in row {header_row}. "
                "Use a header name, Excel letters such as A/AA, or a 1-based number."
            ) from error
    return column_index, header_row + 1 if header_row > 0 else 1


def read_oem_values(input_path: Path, sheet_name: str | None, column: str, header_row: int) -> list[object]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet {sheet_name!r} was not found in {input_path.name}.")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        column_index, first_data_row = resolve_column(sheet, column, header_row)
        return [
            row[0]
            for row in sheet.iter_rows(
                min_row=first_data_row,
                min_col=column_index,
                max_col=column_index,
                values_only=True,
            )
            if row[0] is not None and str(row[0]).strip()
        ]
    finally:
        workbook.close()


def find_search_box(page: Page) -> Locator | None:
    search_box = page.locator(SEARCH_INPUT_SELECTOR)
    if search_box.count() and search_box.is_visible() and search_box.is_enabled():
        return search_box
    return None


def find_search_button(page: Page) -> Locator | None:
    search_button = page.locator(SEARCH_BUTTON_SELECTOR)
    if search_button.count() and search_button.is_visible() and search_button.is_enabled():
        return search_button
    return None


def dismiss_blocking_popups(page: Page, press_escape: bool = True) -> None:
    close_buttons = page.locator(
        'button[aria-label*="close" i], button[title*="close" i], '
        'button[aria-label*="dismiss" i], button:has-text("Close"), '
        'button:has-text("Dismiss"), button:has-text("Not now"), button:has-text("No thanks")'
    )
    for _ in range(20):
        dismissed = False
        for index in range(close_buttons.count()):
            button = close_buttons.nth(index)
            try:
                if button.is_visible() and button.is_enabled():
                    button.click(timeout=1_500)
                    dismissed = True
                    break
            except PlaywrightError:
                continue
        if not dismissed:
            break

    if press_escape:
        try:
            page.keyboard.press("Escape")
        except PlaywrightError:
            pass


def page_needs_human_verification(page: Page) -> bool:
    try:
        title = page.title().casefold()
        body_text = page.locator("body").inner_text(timeout=5_000).casefold()
    except PlaywrightError:
        return False
    return (
        "performing security verification" in body_text
        or "security verification" in body_text
        or "verify you are human" in body_text
        or "checking your browser" in body_text
        or "just a moment" in title
    )


def is_hyundai_site_url(url: str) -> bool:
    return urlparse(url).hostname == urlparse(SITE_URL).hostname


def wait_for_human_verification(page: Page, challenge_failures: list[str]) -> None:
    if not page_needs_human_verification(page):
        return

    print(
        "The site is asking for human verification. Complete the check in the opened browser; "
        "the crawler will continue automatically.",
        flush=True,
    )
    reported_failures = 0
    try:
        while page_needs_human_verification(page):
            while reported_failures < len(challenge_failures):
                print(
                    "A verification request failed: "
                    f"{challenge_failures[reported_failures]}. The browser will remain open; "
                    "complete the check manually when it is available.",
                    flush=True,
                )
                reported_failures += 1
            page.wait_for_timeout(1_000)
    except PlaywrightError as error:
        if page.is_closed():
            raise HumanVerificationError("The browser tab was closed before verification completed.") from error
        raise


def open_site(page: Page, challenge_failures: list[str]) -> None:
    if not is_hyundai_site_url(page.url):
        page.goto(SITE_URL, wait_until="domcontentloaded", timeout=60_000)
    else:
        try:
            page.wait_for_load_state("domcontentloaded", timeout=30_000)
        except PlaywrightTimeoutError:
            pass
    wait_for_human_verification(page, challenge_failures)
    if find_search_box(page) is None:
        page_description = f"Title: {page.title()!r}; URL: {page.url}; "
        body_excerpt = page.locator("body").inner_text(timeout=15_000).strip().replace("\n", " ")[:300]
        raise RuntimeError(
            "The HyundaiPartsDeal search box was not found after page load. "
            f"{page_description}Page text: {body_excerpt!r}"
        )


def find_matching_result(page: Page, oem_number: object) -> Locator | None:
    if not normalize_part_number(oem_number):
        return None

    links = page.locator('a[href*="/genuine/"]')
    for index in range(links.count()):
        link = links.nth(index)
        try:
            if not link.is_visible():
                continue
            text = link.inner_text().strip()
            href = link.get_attribute("href") or ""
        except PlaywrightError:
            continue
        expected = normalize_part_number(oem_number)
        if expected in normalize_part_number(text) or expected in normalize_part_number(href):
            return link
    return None


def wait_for_matching_result(page: Page, oem_number: object, timeout_seconds: float = 12) -> Locator | None:
    deadline = time.monotonic() + timeout_seconds
    while True:
        result_link = find_matching_result(page, oem_number)
        if result_link is not None:
            return result_link
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return None
        page.wait_for_timeout(min(500, int(remaining * 1_000)))


def wait_for_product_suggestion(
    page: Page, oem_number: object, timeout_seconds: float = 3
) -> Locator | None:
    return wait_for_matching_result(page, oem_number, timeout_seconds)


def is_product_detail_page(page: Page) -> bool:
    if not urlparse(page.url).path.startswith("/genuine/"):
        return False
    detail_text = page.locator("body").inner_text(timeout=2_000)
    return bool(
        re.search(r"(?i)Hyundai\s+[A-Z0-9-]+", detail_text)
        or re.search(r"(?im)^Manufacturer Part Number\s*$", detail_text)
    )


def wait_for_product_detail(page: Page, timeout_seconds: float = 20) -> bool:
    deadline = time.monotonic() + timeout_seconds
    while True:
        if is_product_detail_page(page):
            return True
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return False
        page.wait_for_timeout(min(250, max(1, int(remaining * 1_000))))


def wait_for_product_msrp(page: Page, timeout_seconds: float = 20) -> str:
    product = page.locator("body")
    deadline = time.monotonic() + timeout_seconds
    latest_text = ""
    while True:
        if product.count() and product.is_visible():
            latest_text = product.inner_text(timeout=2_000)
            if parse_product_msrp(latest_text) is not None:
                return latest_text
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return latest_text
        page.wait_for_timeout(min(500, max(1, int(remaining * 1_000))))


def is_matching_detail_page(page: Page, oem_number: object) -> bool:
    expected = normalize_part_number(oem_number)
    if not expected:
        return False

    if detail_url_matches_oem(page.url, oem_number):
        return True

    body_text = page.locator("body").inner_text(timeout=15_000)
    part_numbers = re.findall(r"(?i)Hyundai\s+#?([A-Z0-9-]*\d[A-Z0-9-]*)", body_text)
    part_numbers.extend(
        re.findall(r"(?im)^Manufacturer Part Number\s*\n\s*([A-Z0-9-]+)", body_text)
    )
    if any(normalize_part_number(part_number) == expected for part_number in part_numbers):
        return True

    return False


def detail_url_matches_oem(url: str, oem_number: object) -> bool:
    expected = normalize_part_number(oem_number)
    final_path_segment = unquote(urlparse(url).path.rstrip("/").rsplit("/", 1)[-1])
    final_path_segment = re.sub(r"\.html?$", "", final_path_segment, flags=re.IGNORECASE)
    return bool(
        expected
        and urlparse(url).path.startswith("/genuine/")
        and normalize_part_number(final_path_segment).endswith(expected)
    )


def detail_status(body_text: str, oem_number: object, msrp: float | None) -> str:
    if msrp is None:
        return "msrp_not_found"
    part_number_match = re.search(r"(?i)Hyundai\s+#?([A-Z0-9-]*\d[A-Z0-9-]*)", body_text)
    if part_number_match is None:
        part_number_match = re.search(
            r"(?im)^Manufacturer Part Number\s*\n\s*([A-Z0-9-]+)", body_text
        )
    current_part = part_number_match.group(1) if part_number_match else ""
    if current_part and normalize_part_number(current_part) != normalize_part_number(oem_number):
        return f"success_superseded: {current_part}"
    return "success"


def crawl_msrp(
    page: Page,
    oem_number: object,
    challenge_failures: list[str],
) -> tuple[float | None, str]:
    dismiss_blocking_popups(page)
    search_box = find_search_box(page)
    if search_box is None:
        if page_needs_human_verification(page):
            wait_for_human_verification(page, challenge_failures)
            search_box = find_search_box(page)
        if search_box is None:
            raise RuntimeError(
                f"Hyundai search is unavailable on {page.url!r} (page title: {page.title()!r})."
            )

    search_box.fill(str(oem_number).strip())

    print("  Entering the OEM and submitting the search.", flush=True)
    starting_url = page.url
    search_box.press("Enter")
    try:
        page.wait_for_url(
            lambda url: url != starting_url,
            wait_until="domcontentloaded",
            timeout=20_000,
        )
    except PlaywrightTimeoutError:
        search_button = find_search_button(page)
        if search_button is not None:
            print("  Enter did not navigate; clicking the search button.", flush=True)
            search_button.click(timeout=10_000)
    print("  Search submitted; checking the product page and results.", flush=True)
    wait_for_human_verification(page, challenge_failures)

    if not is_product_detail_page(page):
        result_link = wait_for_matching_result(page, oem_number)
        if result_link is None:
            print(f"  No Hyundai product result matched OEM {oem_number!s}.")
            return None, "not_found"

        try:
            result_link.click(timeout=10_000)
            page.wait_for_load_state("domcontentloaded", timeout=30_000)
        except PlaywrightTimeoutError:
            pass

    wait_for_human_verification(page, challenge_failures)
    if not wait_for_product_detail(page):
        print(f"  Product detail did not open after searching {oem_number!s}.")
        return None, f"error: product detail page did not open for {oem_number!s}"
    if not is_matching_detail_page(page, oem_number):
        print(f"  Hyundai product page did not match OEM {oem_number!s}.")
        return None, "not_found"
    dismiss_blocking_popups(page)
    detail_text = wait_for_product_msrp(page)
    price = parse_product_msrp(detail_text)
    if price is None:
        print(f"  Detail page opened, but no MSRP label/value was found for {oem_number!s}.")
    else:
        print(f"  MSRP: ${price:,.2f}")
    status = detail_status(detail_text, oem_number, price)
    return price, status


def save_results(output_path: Path, results: list[tuple[object, float | None, date, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(OUTPUT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for oem_number, msrp, crawled_on, status in results:
        sheet.append((oem_number, msrp, crawled_on, status))
        sheet.cell(sheet.max_row, 2).number_format = '$#,##0.00'
        sheet.cell(sheet.max_row, 3).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 24
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()


def load_results(output_path: Path) -> list[tuple[object, float | None, date, str]]:
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [str(cell or "").strip().casefold() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        indexes = {name: header.index(name) for name in ("oem number", "msrp", "date") if name in header}
        if "oem number" not in indexes:
            return []
        has_status = "status" in header
        status_index = header.index("status") if has_status else -1
        loaded: list[tuple[object, float | None, date, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            oem_number = row[indexes["oem number"]] if indexes["oem number"] < len(row) else None
            if oem_number is None or not str(oem_number).strip():
                continue
            msrp = row[indexes["msrp"]] if "msrp" in indexes and indexes["msrp"] < len(row) else None
            crawled_on = row[indexes["date"]] if "date" in indexes and indexes["date"] < len(row) else date.today()
            status = str(row[status_index] or "").strip() if has_status and status_index < len(row) else ""
            if not status:
                status = "success" if msrp is not None else "error: legacy row has no status or MSRP"
            loaded.append((oem_number, msrp, crawled_on, status))
        return loaded
    finally:
        workbook.close()


def status_is_complete(status: str) -> bool:
    if status.casefold() == "success":
        return True
    replacement_match = re.fullmatch(
        r"success_superseded:\s*([A-Z0-9][A-Z0-9-]*)",
        status.strip(),
        re.IGNORECASE,
    )
    return bool(replacement_match and re.search(r"\d", replacement_match.group(1)))


def prepare_crawl_rows(
    oem_values: list[object],
    previous_results: list[tuple[object, float | None, date, str]],
) -> tuple[list[tuple[object, float | None, date, str]], list[tuple[int, object]]]:
    previous_by_oem: dict[str, list[tuple[object, float | None, date, str]]] = {}
    for row in previous_results:
        previous_by_oem.setdefault(normalize_part_number(row[0]), []).append(row)

    results: list[tuple[object, float | None, date, str]] = []
    pending: list[tuple[int, object]] = []
    for row_index, oem_number in enumerate(oem_values):
        key = normalize_part_number(oem_number)
        prior_rows = previous_by_oem.get(key, [])
        prior = prior_rows.pop(0) if prior_rows else None
        if prior is not None and status_is_complete(prior[3]):
            results.append((oem_number, prior[1], prior[2], prior[3]))
        else:
            if prior is None:
                results.append((oem_number, None, date.today(), "pending"))
            else:
                results.append((oem_number, prior[1], prior[2], prior[3]))
            pending.append((row_index, oem_number))
    return results, pending


def find_chrome(custom_path: Path | None = None) -> Path | None:
    if custom_path is not None:
        candidate = custom_path.expanduser().resolve()
        return candidate if candidate.is_file() else None

    candidates = [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "C:/Users/Default/AppData/Local")) / "Google/Chrome/Application/chrome.exe",
    ]
    return next((candidate for candidate in candidates if candidate.is_file()), None)


def cdp_is_ready() -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{CDP_PORT}/json/version", timeout=1):
            return True
    except Exception:
        return False


def cdp_page_is_hyundai() -> bool:
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{CDP_PORT}/json/list", timeout=2) as response:
            pages = json.loads(response.read().decode("utf-8"))
    except (Exception, ValueError):
        return False

    for page in pages:
        url = str(page.get("url") or "")
        if is_hyundai_site_url(url):
            return True
    return False


def should_start_chrome() -> bool:
    return not (cdp_is_ready() and cdp_page_is_hyundai())


def start_chrome(chrome_path: Path, profile_path: Path) -> subprocess.Popen:
    profile_path.mkdir(parents=True, exist_ok=True)
    process = subprocess.Popen(
        [
            str(chrome_path),
            f"--remote-debugging-port={CDP_PORT}",
            f"--user-data-dir={profile_path}",
            "--no-first-run",
            "--no-default-browser-check",
            SITE_URL,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    for _ in range(60):
        if cdp_is_ready() and cdp_page_is_hyundai():
            return process
        if process.poll() is not None:
            raise RuntimeError(f"Chrome exited during startup with code {process.returncode}.")
        time.sleep(0.5)
    process.terminate()
    process.wait(timeout=5)
    raise RuntimeError(f"Chrome started but the Hyundai page was not available on port {CDP_PORT}.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Crawl Hyundai parts MSRP values from an Excel OEM list.")
    parser.add_argument("--input", required=True, type=Path, help="Input .xlsx or .xlsm workbook.")
    parser.add_argument(
        "--column",
        required=True,
        help="OEM column header, Excel column letters (for example A), or a 1-based column number.",
    )
    parser.add_argument("--output", type=Path, help="Output workbook (default: <input>_results.xlsx).")
    parser.add_argument("--sheet", help="Worksheet name (default: active worksheet).")
    parser.add_argument("--chrome-path", type=Path, help="Path to installed Google Chrome executable.")
    crawl_mode = parser.add_mutually_exclusive_group()
    crawl_mode.add_argument(
        "--resume",
        dest="crawl_mode",
        action="store_const",
        const="resume",
        help="Resume from the existing output workbook (default).",
    )
    crawl_mode.add_argument(
        "--start-over",
        dest="crawl_mode",
        action="store_const",
        const="start-over",
        help="Ignore existing results and crawl every input OEM again.",
    )
    parser.set_defaults(crawl_mode="resume")
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="Header row number; use 0 when the input has no header (default: 1).",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    if not input_path.is_file():
        print(f"Input workbook does not exist: {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        print("Input must be an .xlsx or .xlsm workbook.", file=sys.stderr)
        return 2
    if args.header_row < 0:
        print("--header-row must be 0 or greater.", file=sys.stderr)
        return 2

    output_path = args.output or input_path.with_name(f"{input_path.stem}_results.xlsx")
    output_path = output_path.expanduser().resolve()
    if output_path == input_path:
        print("Input and output paths must be different.", file=sys.stderr)
        return 2

    try:
        oem_values = read_oem_values(input_path, args.sheet, args.column, args.header_row)
    except (OSError, ValueError, KeyError) as error:
        print(f"Could not read input workbook: {error}", file=sys.stderr)
        return 2
    if not oem_values:
        print("No OEM values were found in the selected column.", file=sys.stderr)
        return 2

    results: list[tuple[object, float | None, date, str]] = []
    if args.crawl_mode == "resume" and output_path.is_file():
        try:
            results = load_results(output_path)
        except (OSError, ValueError, KeyError) as error:
            print(f"Could not load existing results for resume: {error}", file=sys.stderr)
            return 2
        if results:
            print(f"Resuming from {output_path}: loaded {len(results)} previous row(s).")

    results, pending_oems = prepare_crawl_rows(oem_values, results)

    if not pending_oems:
        print("All input OEM numbers already have completed statuses. Use --start-over to crawl them again.")
        return 0

    chrome_path = find_chrome(args.chrome_path)
    if chrome_path is None:
        if args.chrome_path:
            print(f"Chrome executable not found: {args.chrome_path}", file=sys.stderr)
        else:
            print("Google Chrome was not found. Install Chrome or pass --chrome-path.", file=sys.stderr)
        return 2

    print(
        f"Loaded {len(oem_values)} OEM number(s); {len(pending_oems)} to crawl "
        f"({args.crawl_mode}). Browser will open for the crawl."
    )
    profile_path = Path(__file__).resolve().parent / ".chrome-profile"
    chrome_process = None
    browser = None
    with sync_playwright() as playwright:
        if should_start_chrome():
            if cdp_is_ready():
                print("Found a stale browser session on the Hyundai CDP port; opening a fresh browser.", flush=True)
                try:
                    with urllib.request.urlopen(f"http://127.0.0.1:{CDP_PORT}/json/list", timeout=2) as response:
                        targets = json.loads(response.read().decode("utf-8"))
                    for target in targets:
                        target_id = target.get("id")
                        if target_id:
                            try:
                                urllib.request.urlopen(
                                    f"http://127.0.0.1:{CDP_PORT}/json/close/{target_id}",
                                    timeout=2,
                                )
                            except Exception:
                                pass
                except Exception:
                    pass
                time.sleep(0.5)
            chrome_process = start_chrome(chrome_path, profile_path)
        try:
            browser = playwright.chromium.connect_over_cdp(f"http://127.0.0.1:{CDP_PORT}")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            pages = context.pages
            page = pages[0] if pages else context.new_page()
            for extra_page in pages[1:]:
                try:
                    extra_page.close()
                except PlaywrightError:
                    pass
            page.set_default_timeout(30_000)

            def record_challenge_failure(request: Request) -> None:
                if "challenges.cloudflare.com" in request.url or "awswaf.com" in request.url:
                    challenge_failures.append(f"{request.url} ({request.failure or 'request failed'})")

            challenge_failures: list[str] = []
            page.on("requestfailed", record_challenge_failure)
            try:
                open_site(page, challenge_failures)
            except HumanVerificationError as error:
                print(f"Crawl stopped: {error}", file=sys.stderr)
                return 1
            for index, (row_index, oem_number) in enumerate(pending_oems, start=1):
                print(f"[{index}/{len(pending_oems)}] Searching {oem_number!s}")
                try:
                    msrp, status = crawl_msrp(
                        page,
                        oem_number,
                        challenge_failures,
                    )
                except HumanVerificationError as error:
                    print(f"Crawl stopped: {error}", file=sys.stderr)
                    return 1
                except Exception as error:
                    print(f"  Crawl failed for {oem_number!s}: {error}")
                    msrp = None
                    status = f"error: {type(error).__name__}: {error}"

                result_row = (oem_number, msrp, date.today(), status)
                results[row_index] = result_row
                save_results(output_path, results)
                dismiss_blocking_popups(page)
                if status_is_complete(status):
                    continue
                if status == "not_found":
                    continue
                print(f"  Status {status!r} is not successful; continuing to the next search.")
        finally:
            if browser is not None:
                try:
                    browser.close()
                except Exception:
                    pass
            if chrome_process.poll() is None:
                chrome_process.terminate()
                try:
                    chrome_process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    chrome_process.kill()

    print(f"Saved {len(results)} row(s) to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
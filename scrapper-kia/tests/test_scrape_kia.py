import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scrape_kia import (
    crawl_msrp,
    detail_status,
    detail_url_matches_oem,
    is_kia_site_url,
    load_results,
    normalize_part_number,
    parse_args,
    parse_msrp,
    parse_product_msrp,
    prepare_crawl_rows,
    read_oem_values,
    resolve_column,
    save_results,
    status_is_complete,
    wait_for_human_verification,
)


class KiaCrawlerTests(unittest.TestCase):
    def test_attempt_on_search_option_is_removed(self):
        with patch("sys.argv", ["scrape_kia.py", "--input", "parts.xlsx", "--column", "A"]):
            args = parse_args()
        self.assertFalse(hasattr(args, "attempt_on_search"))

    def test_missing_suggestion_submits_full_search_once(self):
        page = Mock()
        search_box = Mock()
        page.url = "https://www.kiapartsnow.com/"
        with (
            patch("scrape_kia.dismiss_blocking_popups"),
            patch("scrape_kia.find_search_box", return_value=search_box),
            patch("scrape_kia.wait_for_product_suggestion", return_value=None),
            patch("scrape_kia.wait_for_human_verification"),
            patch("scrape_kia.is_product_detail_page", return_value=False),
            patch("scrape_kia.wait_for_matching_result", return_value=None),
        ):
            result = crawl_msrp(page, "0K2A1-09-000", [])
        self.assertEqual(result, (None, "not_found"))
        search_box.press.assert_called_once_with("Enter")

    def test_verification_waits_for_manual_completion_after_request_failure(self):
        page = Mock()
        failures = ["challenge request failed: net::ERR_NAME_NOT_RESOLVED"]
        with patch("scrape_kia.page_needs_human_verification", side_effect=[True, True, False]):
            wait_for_human_verification(page, failures)
        page.reload.assert_not_called()
        self.assertEqual(page.wait_for_timeout.call_count, 1)

    def test_normalize_part_number(self):
        self.assertEqual(normalize_part_number("0K2A1-09-000"), "0K2A109000")

    def test_parse_msrp_with_currency_and_thousands_separator(self):
        self.assertEqual(parse_msrp("Kia 0K2A1-09-000\nMSRP: $1,030.06"), 1030.06)

    def test_parse_product_price_uses_original_msrp_not_discount(self):
        text = """Kia 0K2A1-09-000 Air Cleaner Filter
Part Description
Filter-Air Cleaner
$23.26 MSRP: $30.06
You Save: $6.80 (23%)
Related Parts
Kia 0K2A1-09-001
MSRP $999.99"""
        self.assertEqual(parse_product_msrp(text), 30.06)

    def test_parse_product_price_falls_back_to_msrp_without_sale_price(self):
        self.assertEqual(parse_product_msrp("Kia 0K2A1-09-000 Air Cleaner\nMSRP $30.06"), 30.06)

    def test_parse_product_price_requires_kia_product_heading(self):
        self.assertIsNone(parse_product_msrp("MSRP $30.06"))

    def test_detail_url_matches_kia_genuine_part_url(self):
        url = "https://www.kiapartsnow.com/genuine/kia-air-cleaner~0k2a1-09-000.html"
        self.assertTrue(detail_url_matches_oem(url, "0K2A1-09-000"))
        self.assertFalse(detail_url_matches_oem(url, "0K2A1-09-001"))

    def test_site_url_recognizes_kia_home_and_product_pages(self):
        self.assertTrue(is_kia_site_url("https://www.kiapartsnow.com/"))
        self.assertTrue(
            is_kia_site_url(
                "https://www.kiapartsnow.com/genuine/kia-air-cleaner~0k2a1-09-000.html"
            )
        )
        self.assertFalse(is_kia_site_url("about:blank"))

    def test_status_reports_superseded_kia_part(self):
        self.assertEqual(
            detail_status("Kia 0K2A1-09-001 Air Cleaner Filter", "0K2A1-09-000", 30.06),
            "success_superseded: 0K2A1-09-001",
        )
        self.assertEqual(detail_status("Kia Parts\nMSRP $30.06", "0K2A1-09-000", 30.06), "success")

    def test_only_successful_statuses_are_complete(self):
        self.assertTrue(status_is_complete("success"))
        self.assertTrue(status_is_complete("success_superseded: 0K2A1-09-001"))
        self.assertFalse(status_is_complete("not_found"))
        self.assertFalse(status_is_complete("msrp_not_found"))
        self.assertFalse(status_is_complete("error: TimeoutError"))

    def test_resolve_column_supports_header_and_excel_letter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("Description", "OEM Number"))
        sheet.append(("Air filter", "0K2A1-09-000"))
        self.assertEqual(resolve_column(sheet, "OEM Number", 1), (2, 2))
        self.assertEqual(resolve_column(sheet, "B", 1), (2, 2))
        workbook.close()

    def test_read_oem_values_from_workbook(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "parts.xlsx"
            workbook = Workbook()
            workbook.active.append(("Description", "OEM Number"))
            workbook.active.append(("Air filter", "0K2A1-09-000"))
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_oem_values(path, None, "OEM Number", 1), ["0K2A1-09-000"])

    def test_duplicate_oems_are_queued_as_separate_rows(self):
        results, pending = prepare_crawl_rows(["0K2A1-09-000", "0K2A1-09-000"], [])
        self.assertEqual(len(results), 2)
        self.assertEqual(pending, [(0, "0K2A1-09-000"), (1, "0K2A1-09-000")])

    def test_resume_matches_duplicate_results_by_occurrence(self):
        saved = [
            ("0K2A1-09-000", 30.06, date(2026, 10, 6), "success"),
            ("0K2A1-09-000", None, date(2026, 10, 6), "not_found"),
        ]
        results, pending = prepare_crawl_rows(["0K2A1-09-000", "0K2A1-09-000"], saved)
        self.assertEqual(results[0][1:], saved[0][1:])
        self.assertEqual(pending, [(1, "0K2A1-09-000")])

    def test_save_and_load_results_round_trip(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            rows = [("0K2A1-09-000", 23.26, date(2026, 10, 6), "success")]
            save_results(output, rows)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                tuple(cell.value for cell in workbook.active[1]),
                ("OEM Number", "MSRP", "Date", "Status"),
            )
            workbook.close()
            self.assertEqual(load_results(output)[0][0:2], ("0K2A1-09-000", 23.26))


if __name__ == "__main__":
    unittest.main()